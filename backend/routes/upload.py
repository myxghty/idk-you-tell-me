import os
from flask import Blueprint, request, jsonify

from utils.ai import run_generation

upload_bp = Blueprint("upload", __name__)

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".pptx"}


def extract_text(file_storage, ext):
    """Extract raw text from an uploaded PDF, DOCX, or PPTX file."""
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(file_storage)
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".docx":
        import docx
        document = docx.Document(file_storage)
        return "\n".join(p.text for p in document.paragraphs)

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(file_storage)
        chunks = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    chunks.append(shape.text_frame.text)
        return "\n".join(chunks)

    return ""


@upload_bp.route("/", methods=["POST"], strict_slashes=False)
def upload():
    try:
        if "file" not in request.files:
            return jsonify({"error": "No 'file' part in the request."}), 400

        file = request.files["file"]
        if not file or file.filename == "":
            return jsonify({"error": "No file selected."}), 400

        filename = file.filename
        ext = os.path.splitext(filename)[1].lower()

        if ext not in ALLOWED_EXTENSIONS:
            return jsonify({
                "error": f"Unsupported file type '{ext}'. Allowed: .pdf, .docx, .pptx"
            }), 400

        level = request.form.get("level")
        language = request.form.get("language")

        try:
            text = extract_text(file, ext)
        except Exception as err:
            print("Text extraction failed:", err)
            return jsonify({"error": f"Could not read {ext} file: {err}"}), 400

        if not text or len(text.strip()) < 20:
            return jsonify({
                "error": "Could not extract enough readable text from this file. "
                         "It may be scanned/image-based or empty."
            }), 400

        data, error = run_generation(text, language=language, level=level)
        if error:
            return jsonify(error), error.get("status", 502)

        data["extracted_text"] = text
        data["filename"] = filename
        return jsonify(data)

    except Exception as err:
        print("Unexpected error in /api/upload:", err)
        return jsonify({"error": "Internal server error"}), 500
